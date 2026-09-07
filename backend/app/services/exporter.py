import io
import re
from typing import List, Dict, Any, Tuple

# ReportLab Imports for PDF Generation
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.pdfgen import canvas

# python-docx Imports for Word Document Generation
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ---------------------------------------------------------------------------
# 1. Screenplay Element Classification Parser
# ---------------------------------------------------------------------------

SCENE_HEADING_PATTERN = re.compile(
    r'^(?:INT|EXT|EST|I/E|INT/EXT|INT\./EXT\.)[\.\s]', re.IGNORECASE
)
TRANSITION_PATTERN = re.compile(
    r'^(?:FADE IN:|FADE IN\.|FADE IN|FADE OUT\.|FADE OUT|CUT TO:|DISSOLVE TO:|MATCH CUT TO:|SMASH CUT TO:|FADE TO BLACK\.|FADE TO BLACK|INTERCUT WITH:).*$|.*TO:$',
    re.IGNORECASE
)
SCENE_TIME_PATTERN = re.compile(
    r'.*\-\s*(?:DAY|NIGHT|EVENING|MORNING|CONTINUOUS|LATER|MOMENTS LATER|SAME TIME|SUNSET|DAWN|DUSK|AFTERNOON)$',
    re.IGNORECASE
)

LOCATION_KEYWORDS = {
    "ROOM", "HOUSE", "BEDROOM", "KITCHEN", "OFFICE", "BASEMENT", "HALLWAY", 
    "STREET", "ALLEY", "CAR", "BUILDING", "APARTMENT", "ROOFTOP", "WAREHOUSE", 
    "HOSPITAL", "COURT", "COURTROOM", "RESTAURANT", "BAR", "PARK", "LOCATION", 
    "STATION", "LAB", "LABORATORY", "CORRIDOR", "CABIN", "GARAGE", "CLUB"
}

CHARACTER_ABBREV = {"O.S.", "V.O.", "O.C.", "CONT'D", "O.S", "V.O", "O.C"}

def is_location_heading(text: str) -> bool:
    """
    Checks if a line represents a scene/location heading (even if missing INT./EXT. prefix).
    """
    upper = text.upper()
    if SCENE_HEADING_PATTERN.match(upper):
        return True
    if SCENE_TIME_PATTERN.match(upper):
        return True
    
    # Standalone location keyword match without INT./EXT. prefix:
    # Must be uppercase in original text, short (<= 4 words, <= 35 chars),
    # and contain a known location keyword.
    if text.isupper() and len(text) <= 35:
        words = set(re.findall(r'\b[A-Z]+\b', upper))
        if words.intersection(LOCATION_KEYWORDS):
            if len(words) <= 4:
                return True
    return False


def classify_screenplay_lines(raw_text: str) -> List[Tuple[str, str]]:
    """
    Parses screenplay text line by line and classifies each line into:
    'SCENE_HEADING', 'CHARACTER', 'PARENTHETICAL', 'DIALOGUE', 'TRANSITION', 'ACTION', 'BLANK'
    """
    lines = raw_text.splitlines()
    classified: List[Tuple[str, str]] = []
    
    prev_type = "BLANK"
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            classified.append(("BLANK", ""))
            prev_type = "BLANK"
            continue

        # 1. Scene Heading / Location check
        if is_location_heading(stripped):
            classified.append(("SCENE_HEADING", stripped.upper()))
            prev_type = "SCENE_HEADING"
            continue

        # 2. Transition check
        if TRANSITION_PATTERN.match(stripped):
            classified.append(("TRANSITION", stripped.upper()))
            prev_type = "TRANSITION"
            continue

        # 3. Parenthetical check (ONLY if directly following a CHARACTER or another PARENTHETICAL)
        if stripped.startswith("(") and stripped.endswith(")"):
            if prev_type in ("CHARACTER", "PARENTHETICAL"):
                classified.append(("PARENTHETICAL", stripped))
                prev_type = "PARENTHETICAL"
            else:
                classified.append(("ACTION", stripped))
                prev_type = "ACTION"
            continue

        # 4. Character Name check
        # A line ending with ')' is only a character if it ends with a valid modifier e.g. (O.S.), (V.O.), (CONT'D)
        ends_with_paren = stripped.endswith(")")
        valid_paren = ends_with_paren and (
            stripped.endswith("(O.S.)") or stripped.endswith("(V.O.)") or 
            stripped.endswith("(O.C.)") or stripped.endswith("(CONT'D)") or 
            stripped.endswith("(O.S)") or stripped.endswith("(V.O)")
        )
        
        is_char_abbrev = stripped.upper() in CHARACTER_ABBREV or valid_paren
        valid_end = (not stripped.endswith(".")) or is_char_abbrev

        if any(stripped.endswith(ch) for ch in ("!", "?", ",", ";", ":")):
            valid_end = False

        if ends_with_paren and not valid_paren:
            is_character = False
        else:
            is_character = (
                stripped.isupper()
                and len(stripped) < 45
                and valid_end
                and prev_type != "CHARACTER"
            )

        if is_character:
            classified.append(("CHARACTER", stripped))
            prev_type = "CHARACTER"
            continue

        # 5. Dialogue check (follows CHARACTER or PARENTHETICAL or DIALOGUE)
        if prev_type in ("CHARACTER", "PARENTHETICAL", "DIALOGUE"):
            classified.append(("DIALOGUE", stripped))
            prev_type = "DIALOGUE"
            continue

        # 6. Default to Action description
        classified.append(("ACTION", stripped))
        prev_type = "ACTION"

    return classified


# ---------------------------------------------------------------------------
# 2. PDF Exporter with Hollywood Standard Layout & Typography
# ---------------------------------------------------------------------------

class NumberedCanvas(canvas.Canvas):
    """
    Canvas for ReportLab to add page numbers in top-right header (starting from page 2).
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            if self._pageNumber > 1:
                self.setFont("Courier", 10)
                # 0.5 inch from top, 1.0 inch from right margin
                self.drawRightString(8.5 * inch - 1.0 * inch, 11.0 * inch - 0.5 * inch, f"{self._pageNumber}.")
            super().showPage()
        super().save()


def generate_pdf_export(
    scenes_text: List[str],
    font_family: str = "Courier",
    font_size: int = 12
) -> bytes:
    """
    Generates a PDF byte stream formatted according to exact Hollywood Screenplay Standards.
    Prevents orphaned character headers using KeepTogether.
    """
    buffer = io.BytesIO()
    
    # 8.5 x 11 inches, Left Margin: 1.5", Right: 1.0", Top: 1.0", Bottom: 1.0"
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * inch,
        rightMargin=1.0 * inch,
        topMargin=1.0 * inch,
        bottomMargin=1.0 * inch
    )

    valid_fonts = {"Courier", "Helvetica", "Times-Roman"}
    pdf_font = font_family if font_family in valid_fonts else "Courier"

    base_size = max(8, min(16, font_size))
    leading = base_size * 1.2

    styles = getSampleStyleSheet()
    
    # 1. Scene Heading: Flush left (0.0"), 12pt Courier UPPERCASE
    style_heading = ParagraphStyle(
        'ScreenplayHeading',
        fontName=pdf_font,
        fontSize=base_size,
        leading=leading,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    # 2. Action: Flush left (0.0"), full width
    style_action = ParagraphStyle(
        'ScreenplayAction',
        fontName=pdf_font,
        fontSize=base_size,
        leading=leading,
        spaceBefore=4,
        spaceAfter=4
    )

    # 3. Character Name: 3.7" from page edge (2.2" from left margin), UPPERCASE
    style_character = ParagraphStyle(
        'ScreenplayCharacter',
        fontName=pdf_font,
        fontSize=base_size,
        leading=leading,
        leftIndent=2.2 * inch,
        spaceBefore=10,
        spaceAfter=0,
        keepWithNext=True
    )

    # 4. Parenthetical: 3.1" from page edge (1.6" from left margin)
    style_parenthetical = ParagraphStyle(
        'ScreenplayParenthetical',
        fontName=pdf_font,
        fontSize=base_size,
        leading=leading,
        leftIndent=1.6 * inch,
        rightIndent=2.0 * inch,
        spaceBefore=0,
        spaceAfter=0,
        keepWithNext=True
    )

    # 5. Dialogue: 2.5" from page edge (1.0" from left margin, 1.5" right indent)
    style_dialogue = ParagraphStyle(
        'ScreenplayDialogue',
        fontName=pdf_font,
        fontSize=base_size,
        leading=leading,
        leftIndent=1.0 * inch,
        rightIndent=1.5 * inch,
        spaceBefore=0,
        spaceAfter=6
    )

    # 6. Transition: Right-aligned bold transition (FADE OUT. / CUT TO:)
    bold_font = f"{pdf_font}-Bold" if pdf_font in ("Courier", "Helvetica", "Times-Roman") else pdf_font
    style_transition = ParagraphStyle(
        'ScreenplayTransition',
        fontName=bold_font,
        fontSize=base_size,
        leading=leading,
        alignment=TA_RIGHT,
        spaceBefore=12,
        spaceAfter=12
    )

    story = []

    for scene_idx, scene_text in enumerate(scenes_text):
        classified = classify_screenplay_lines(scene_text)
        
        i = 0
        while i < len(classified):
            item_type, text = classified[i]

            if item_type == "BLANK":
                story.append(Spacer(1, leading))
                i += 1
                continue

            if item_type == "SCENE_HEADING":
                story.append(Paragraph(text, style_heading))
                i += 1
                continue

            if item_type == "ACTION":
                action_lines = []
                while i < len(classified) and classified[i][0] == "ACTION":
                    action_lines.append(classified[i][1])
                    i += 1
                story.append(Paragraph(" ".join(action_lines), style_action))
                continue

            if item_type == "TRANSITION":
                story.append(Paragraph(text, style_transition))
                i += 1
                continue

            # Orphan Prevention: Group CHARACTER + PARENTHETICAL + DIALOGUE together
            if item_type == "CHARACTER":
                block_elements = []
                block_elements.append(Paragraph(text, style_character))
                i += 1

                current_dialogue_lines = []

                while i < len(classified) and classified[i][0] in ("PARENTHETICAL", "DIALOGUE"):
                    sub_type, sub_text = classified[i]
                    if sub_type == "PARENTHETICAL":
                        if current_dialogue_lines:
                            block_elements.append(Paragraph(" ".join(current_dialogue_lines), style_dialogue))
                            current_dialogue_lines = []
                        block_elements.append(Paragraph(sub_text, style_parenthetical))
                    elif sub_type == "DIALOGUE":
                        current_dialogue_lines.append(sub_text)
                    i += 1

                if current_dialogue_lines:
                    block_elements.append(Paragraph(" ".join(current_dialogue_lines), style_dialogue))

                story.append(KeepTogether(block_elements))
                continue

            # Fallback for standalone dialogue lines
            if item_type == "PARENTHETICAL":
                story.append(Paragraph(text, style_parenthetical))
                i += 1
            elif item_type == "DIALOGUE":
                dialogue_lines = []
                while i < len(classified) and classified[i][0] == "DIALOGUE":
                    dialogue_lines.append(classified[i][1])
                    i += 1
                story.append(Paragraph(" ".join(dialogue_lines), style_dialogue))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 3. DOCX Exporter with Industry Layout & Typography
# ---------------------------------------------------------------------------

def generate_docx_export(
    scenes_text: List[str],
    font_family: str = "Courier New",
    font_size: int = 12
) -> bytes:
    """
    Generates a Microsoft Word (.docx) byte stream formatted according to screenplay standards.
    """
    doc = docx.Document()

    # Set 1.5" Left, 1.0" Right/Top/Bottom Margins
    for section in doc.sections:
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    base_font = font_family or "Courier New"
    pt_size = Pt(font_size)

    def add_styled_paragraph(text: str, line_type: str):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = base_font
        run.font.size = pt_size

        if line_type == "SCENE_HEADING":
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
        elif line_type == "CHARACTER":
            p.paragraph_format.left_indent = Inches(2.2)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
        elif line_type == "PARENTHETICAL":
            p.paragraph_format.left_indent = Inches(1.6)
            p.paragraph_format.right_indent = Inches(2.0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
        elif line_type == "DIALOGUE":
            p.paragraph_format.left_indent = Inches(1.0)
            p.paragraph_format.right_indent = Inches(1.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
        elif line_type == "TRANSITION":
            run.font.bold = True
            upper_text = text.upper()
            if "FADE IN" in upper_text:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
        else: # ACTION
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)

    for scene_text in scenes_text:
        classified = classify_screenplay_lines(scene_text)
        i = 0
        while i < len(classified):
            item_type, text = classified[i]

            if item_type == "BLANK":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                i += 1
            elif item_type == "ACTION":
                action_lines = []
                while i < len(classified) and classified[i][0] == "ACTION":
                    action_lines.append(classified[i][1])
                    i += 1
                add_styled_paragraph(" ".join(action_lines), "ACTION")
            elif item_type == "CHARACTER":
                add_styled_paragraph(text, "CHARACTER")
                i += 1
                current_dialogue_lines = []

                while i < len(classified) and classified[i][0] in ("PARENTHETICAL", "DIALOGUE"):
                    sub_type, sub_text = classified[i]
                    if sub_type == "PARENTHETICAL":
                        if current_dialogue_lines:
                            add_styled_paragraph(" ".join(current_dialogue_lines), "DIALOGUE")
                            current_dialogue_lines = []
                        add_styled_paragraph(sub_text, "PARENTHETICAL")
                    elif sub_type == "DIALOGUE":
                        current_dialogue_lines.append(sub_text)
                    i += 1

                if current_dialogue_lines:
                    add_styled_paragraph(" ".join(current_dialogue_lines), "DIALOGUE")
            else:
                add_styled_paragraph(text, item_type)
                i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 4. Fountain Exporter & Importer
# ---------------------------------------------------------------------------

def generate_fountain_export(scenes_text: List[str]) -> str:
    """
    Generates a standard Fountain plain text screenplay document string.
    """
    fountain_blocks = []
    for idx, text in enumerate(scenes_text):
        cleaned = text.strip()
        if cleaned:
            fountain_blocks.append(cleaned)
    return "\n\n".join(fountain_blocks) + "\n"


def parse_fountain_import(fountain_content: str) -> List[Dict[str, Any]]:
    """
    Parses an incoming .fountain or plain text script file into scene dict objects.
    Splits scenes based on standard screenplay scene headings (INT, EXT, EST, etc.),
    ignoring title page metadata preambles before the first scene heading.
    """
    lines = fountain_content.splitlines()
    scenes = []
    current_scene_lines = []
    scene_counter = 1
    found_first_heading = False

    for line in lines:
        stripped = line.strip()
        if is_location_heading(stripped):
            if not found_first_heading:
                found_first_heading = True
                current_scene_lines = [line]
                continue
            else:
                if current_scene_lines:
                    raw = "\n".join(current_scene_lines).strip()
                    if raw:
                        scenes.append({
                            "scene_number": scene_counter,
                            "raw_text": raw
                        })
                        scene_counter += 1
                    current_scene_lines = []

        if found_first_heading:
            current_scene_lines.append(line)

    if found_first_heading and current_scene_lines:
        raw = "\n".join(current_scene_lines).strip()
        if raw:
            scenes.append({
                "scene_number": scene_counter,
                "raw_text": raw
            })

    return scenes
